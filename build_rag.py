"""
build_rag.py  —  第一步：建向量库
用法:
    python build_rag.py            # 扫描 DB_DIR，建库保存到 FAISS_DIR
    python build_rag.py --rebuild  # 强制重建（即使库已存在）

建好后 FAISS_DIR 里会有 index.faiss + index.pkl，供 agent.py 加载。
"""

import os, sys, glob, argparse, pickle
from pathlib import Path

import torch
import openpyxl
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document

# ── 配置 ─────────────────────────────────────────────────────────────────────
DB_DIR    = "/gz-data/database"          # 镜头数据根目录
FAISS_DIR = "/gz-data/faiss_index"       # 向量库输出目录
EMB_MODEL = "shibing624/text2vec-base-chinese"

# optical_calculator 路径
sys.path.insert(0, "/gz-data")

# ── Embedding 模型（自动选 GPU）───────────────────────────────────────────────
_device = "cuda" if torch.cuda.is_available() else "cpu"
print(f"Embedding 设备: {_device}")
embeddings = HuggingFaceEmbeddings(
    model_name=EMB_MODEL,
    model_kwargs={"device": _device},
    encode_kwargs={"batch_size": 256},
)

# ── optical_calculator ────────────────────────────────────────────────────────
try:
    from optical_calculator import paraxial_trace, lens_to_surfaces
    _HAS_OPTICS = True
    print("✅ optical_calculator 加载成功")
except ImportError:
    _HAS_OPTICS = False
    print("⚠ optical_calculator 未找到，将跳过近轴预计算")

# ── 解析函数 ──────────────────────────────────────────────────────────────────
def _sf(v, default=None):
    try:
        f = float(str(v).replace("'", "").strip())
        return None if f != f else f
    except:
        return default

def _parse_meta(s, prefix):
    try:
        return float(str(s).replace(prefix, "").strip())
    except:
        return None

def _parse_excel(path: str) -> list:
    """
    解析单个 xlsx 文件。
    文件结构：
        行A: num_X  fov_35  fnum_2.9  aper_12  wave_[...]  fit_0.046
        行B: #  Type  Comment  Radius  Thickness  Material ...  ← 列头，跳过
        行C+: 面数据 (STANDARD 行)
    """
    wb_rows = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            wb_rows.append(list(row))
        wb.close()
    except Exception as e:
        print(f"  ⚠ 读取失败 {path}: {e}")
        return []

    lenses, i = [], 0
    while i < len(wb_rows):
        row = wb_rows[i]
        if str(row[0]).startswith("num_"):
            lens = {
                "source": path,
                "fov":    _parse_meta(row[1], "fov_"),
                "fnum":   _parse_meta(row[2], "fnum_"),
                "aper":   _parse_meta(row[3], "aper_"),
                "fit":    _parse_meta(row[5], "fit_"),   # Zemax merit function
                "surfaces": [],
            }
            j = i + 2   # 跳过列头行
            while j < len(wb_rows):
                r = wb_rows[j]
                if str(r[0]).startswith("num_"):
                    break
                if str(r[1]) == "STANDARD":
                    mat = str(r[5]).strip() if len(r) > 5 else ""
                    lens["surfaces"].append({
                        "surface_num":   _sf(r[0]),
                        "radius":        _sf(r[3]),
                        "thickness":     _sf(r[4], 0.0),
                        "material":      mat if mat not in ("0", "", "nan") else "AIR",
                        "semi_diameter": _sf(r[7], 0.0) if len(r) > 7 else 0.0,
                    })
                j += 1
            lenses.append(lens)
            i = j
        else:
            i += 1
    return lenses

def _parse_excel_safe(path: str) -> list:
    """多进程 worker，捕获异常防止进程崩溃"""
    try:
        return _parse_excel(path)
    except Exception as e:
        return []   # 静默跳过，主进程不打印（避免多进程 print 乱序）

def load_all_lenses(db_dir: str) -> list:
    from multiprocessing import Pool, cpu_count

    patterns = ["**/read_data.xlsx", "**/data.xlsx"]
    files = []
    for pat in patterns:
        files.extend(glob.glob(f"{db_dir}/{pat}", recursive=True))
    files = sorted(set(files))
    if not files:
        files = sorted(glob.glob(f"{db_dir}/**/*.xlsx", recursive=True))
    print(f"找到 {len(files)} 个文件")

    # 用所有 CPU 核并行解析，IO 密集型用 2x 核数
    n_workers = min(cpu_count() * 2, 32)
    print(f"并行解析，workers={n_workers}…")

    lenses = []
    with Pool(n_workers) as pool:
        results = pool.map(_parse_excel_safe, files)
    for r in results:
        lenses.extend(r)

    print(f"共解析 {len(lenses)} 条镜头")
    return lenses

# ── 近轴预计算 ────────────────────────────────────────────────────────────────
def _calc_lens(lens: dict) -> dict:
    if not _HAS_OPTICS:
        return {"valid": False, "msg": "optical_calculator 未加载"}
    try:
        surfs = lens_to_surfaces(lens)
        return paraxial_trace(surfs, lens_meta=lens)
    except Exception as e:
        return {"valid": False, "msg": str(e)}

def _calc_one(lens: dict) -> dict:
    """多进程 worker：追迹单条镜头，返回结果 dict"""
    r = _calc_lens(lens)
    if r.get("valid"):
        return {
            "calc_effl": r["effl"],
            "calc_totr": r["totr"],
            "calc_rms":  r["rms"],
            "calc_yimg": r["y_image"],
            "calc_aper": r["aperture_radius"],
        }
    return {k: None for k in ("calc_effl","calc_totr","calc_rms","calc_yimg","calc_aper")}

def precompute(lenses: list) -> list:
    from multiprocessing import Pool, cpu_count
    n_workers = min(cpu_count(), 16)
    print(f"近轴追迹预计算，workers={n_workers}…")

    with Pool(n_workers) as pool:
        results = pool.map(_calc_one, lenses)

    ok = fail = 0
    for lens, res in zip(lenses, results):
        lens.update(res)
        if lens.get("calc_rms") is not None:
            ok += 1
        else:
            fail += 1
    print(f"  追迹完成：{ok} 成功 / {fail} 失败")
    return lenses

# ── 向量文本生成 ──────────────────────────────────────────────────────────────
def _rms_label(rms):
    if rms is None:   return "未知"
    if rms < 0.05:    return "极小(优秀)"
    if rms < 0.15:    return "小(良好)"
    if rms < 0.5:     return "中等"
    return "大(较差)"

def lens_to_doc(lens: dict, idx: int) -> Document:
    surfs  = lens["surfaces"]
    mats   = list(dict.fromkeys(s["material"] for s in surfs if s["material"] != "AIR"))
    n_ele  = sum(1 for s in surfs if s["material"] != "AIR")
    rms    = lens.get("calc_rms")
    effl   = lens.get("calc_effl")
    rms_s  = f"{rms:.4f}"  if rms  is not None else "未知"
    effl_s = f"{effl:.2f}" if effl is not None else "未知"

    text = (
        f"镜头#{idx} "
        f"视场角FOV={lens.get('fov')}度 全视场{lens.get('fov')}度 "
        f"光圈F/{lens.get('fnum')} F数{lens.get('fnum')} "
        f"入瞳直径{lens.get('aper')}mm 通光口径{lens.get('aper')}mm "
        f"EFFL={effl_s}mm 焦距{effl_s}mm "
        f"近轴RMS={rms_s}mm RMS{_rms_label(rms)} "
        f"镜片数={n_ele}片 玻璃={','.join(mats)} "
        f"来源={Path(lens.get('source', '')).name}"
    )
    return Document(
        page_content=text,
        metadata={
            **{k: v for k, v in lens.items() if k != "surfaces"},
            "surfaces":  str(surfs),
            "lens_idx":  idx,
        },
    )

# ── 建库 ──────────────────────────────────────────────────────────────────────
def build(lenses: list):
    docs = [lens_to_doc(l, i) for i, l in enumerate(lenses)]
    print(f"向量化 {len(docs)} 条文档…")
    vs = FAISS.from_documents(docs, embeddings)
    Path(FAISS_DIR).mkdir(parents=True, exist_ok=True)
    vs.save_local(FAISS_DIR)

    # 同时把 lenses list（含预计算结果）pickle 存下来，供 agent 直接加载
    pkl_path = str(Path(FAISS_DIR) / "lenses.pkl")
    with open(pkl_path, "wb") as f:
        pickle.dump(lenses, f)

    print(f"✅ 向量库已保存: {FAISS_DIR}")
    print(f"✅ 镜头数据已保存: {pkl_path}")
    return vs

# ── 入口 ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--rebuild", action="store_true", help="强制重建（忽略已有库）")
    args = parser.parse_args()

    faiss_exists = (Path(FAISS_DIR) / "index.faiss").exists()
    if faiss_exists and not args.rebuild:
        print(f"向量库已存在: {FAISS_DIR}（跳过，用 --rebuild 强制重建）")
        sys.exit(0)

    lenses = load_all_lenses(DB_DIR)
    lenses = precompute(lenses)
    build(lenses)
